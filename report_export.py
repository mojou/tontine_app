# ============================================================
# RAPPORTS : construction des données + export PDF / CSV / Excel
# ============================================================
# build_report() prépare un rapport (titre, colonnes, lignes, totaux, synthèse)
# à partir des données de la tontine courante (filtrage multi-tenant automatique).
# Les fonctions to_pdf / to_csv / to_excel le mettent en forme.

import csv
import io
from datetime import date, datetime
from decimal import Decimal

import finance
from models import Loan, Sanction, Transaction

REPORT_TYPES = {
    'cotisations': 'Cotisations et tontines',
    'loans': 'Emprunts',
    'sanctions': 'Sanctions',
    'financial': 'Rapport financier (entrées / sorties)',
}

# Nom court en français (onglet Excel, nom du fichier) : jamais de code interne ni d'anglais
REPORT_SHORT_NAMES = {
    'cotisations': 'Cotisations',
    'loans': 'Emprunts',
    'sanctions': 'Sanctions',
    'financial': 'Financier',
}


def report_filename(kind, start, end):
    """Ex. « Rapport Emprunts du 01-01-2026 au 31-12-2026 » (sans extension)"""
    return f"Rapport {REPORT_SHORT_NAMES.get(kind, 'Ghelia')} du {start:%d-%m-%Y} au {end:%d-%m-%Y}"


# Couleurs Ghelia Finance
GREEN = '#14532d'
GREEN_LIGHT = '#e6f1e9'
GOLD = '#c9a227'
GOLD_LIGHT = '#faf2d3'


class Report:
    def __init__(self, kind, title, start, end, columns, money_columns=()):
        self.kind = kind
        self.title = title
        self.start = start
        self.end = end
        self.columns = columns
        self.money_columns = set(money_columns)   # index des colonnes en FCFA
        self.rows = []
        self.totals = None                        # ligne de total (même nombre de colonnes)
        self.summary = []                         # [(libellé, montant)]

    @property
    def period(self):
        return f"du {self.start.strftime('%d/%m/%Y')} au {self.end.strftime('%d/%m/%Y')}"


def _member(obj):
    return obj.member.full_name if getattr(obj, 'member', None) else 'Inconnu'


def _d(value):
    return Decimal(str(value or 0))


def build_report(kind, start, end):
    if kind not in REPORT_TYPES:
        raise ValueError('Type de rapport inconnu')
    title = REPORT_TYPES[kind]
    in_period = [Transaction.date >= start, Transaction.date <= end]

    if kind == 'cotisations':
        report = Report(kind, title, start, end, ['Date', 'Membre', 'Rubrique', 'Mode', 'Référence', 'Montant (FCFA)'], [5])
        txs = (Transaction.query.filter(*in_period, Transaction.type.in_(finance.CONTRIBUTION_CATEGORIES))
               .order_by(Transaction.date, Transaction.id).all())
        by_type = {}
        for t in txs:
            report.rows.append([t.date, _member(t), t.get_type_display(), finance.PAYMENT_MODES.get(t.payment_mode, t.payment_mode or ''),
                                t.payment_reference or '', _d(t.amount)])
            by_type[t.type] = by_type.get(t.type, Decimal('0')) + _d(t.amount)
        total = sum((r[5] for r in report.rows), Decimal('0'))
        report.totals = ['TOTAL', f'{len(report.rows)} paiement(s)', '', '', '', total]
        report.summary = [(finance.type_label(k), v) for k, v in sorted(by_type.items(), key=lambda kv: -kv[1])]
        return report

    if kind == 'loans':
        report = Report(kind, title, start, end,
                        ['Demande', 'Membre', 'Montant', 'Intérêts', 'Total dû', 'Remboursé', 'Reste', 'Échéance', 'Statut'],
                        [2, 3, 4, 5, 6])
        loans = Loan.query.filter(Loan.request_date >= start, Loan.request_date <= end).order_by(Loan.request_date, Loan.id).all()
        for l in loans:
            report.rows.append([l.request_date, _member(l), _d(l.amount), _d(l.interest), _d(l.total_amount),
                                _d(l.amount_paid), _d(l.remaining_amount), l.end_date, l.get_status_display()])
        sums = [sum((r[i] for r in report.rows), Decimal('0')) for i in (2, 3, 4, 5, 6)]
        report.totals = ['TOTAL', f'{len(report.rows)} emprunt(s)', *sums, '', '']
        report.summary = [('Prêté', sums[0]), ('Intérêts attendus', sums[1]), ('Déjà remboursé', sums[3]), ('Reste à recouvrer', sums[4])]
        return report

    if kind == 'sanctions':
        report = Report(kind, title, start, end, ['Date', 'Membre', 'Motif', 'Description', 'Montant (FCFA)', 'Statut'], [4])
        sanctions = (Sanction.query.filter(Sanction.sanction_date >= start, Sanction.sanction_date <= end)
                     .order_by(Sanction.sanction_date, Sanction.id).all())
        paid = pending = Decimal('0')
        for s in sanctions:
            report.rows.append([s.sanction_date, _member(s), s.get_type_display(), (s.description or '')[:80], _d(s.amount),
                                'Payée' if s.status == 'PAID' else 'En attente'])
            if s.status == 'PAID':
                paid += _d(s.amount)
            else:
                pending += _d(s.amount)
        report.totals = ['TOTAL', f'{len(report.rows)} sanction(s)', '', '', paid + pending, '']
        report.summary = [('Payées', paid), ('En attente', pending)]
        return report

    # Rapport financier : toutes les opérations, entrées et sorties séparées
    report = Report(kind, title, start, end, ['Date', 'Membre', 'Opération', 'Fonds', 'Entrée (FCFA)', 'Sortie (FCFA)'], [4, 5])
    txs = Transaction.query.filter(*in_period).order_by(Transaction.date, Transaction.id).all()
    total_in = total_out = Decimal('0')
    by_type = {}
    for t in txs:
        amount = _d(t.amount)
        out = finance.is_outflow(t.type)
        report.rows.append([t.date, _member(t), t.get_type_display(), finance.FUNDS.get(finance.fund_of(t.type), ''),
                            None if out else amount, amount if out else None])
        if out:
            total_out += amount
        else:
            total_in += amount
        by_type[t.type] = by_type.get(t.type, Decimal('0')) + (-amount if out else amount)
    report.totals = ['TOTAL', f'{len(report.rows)} opération(s)', '', '', total_in, total_out]
    report.summary = [('Total des entrées', total_in), ('Total des sorties', total_out), ('Solde de la période', total_in - total_out)]
    report.summary += [(f"  {finance.type_label(k)}", v) for k, v in sorted(by_type.items(), key=lambda kv: -abs(kv[1]))]
    return report


# ---------------------------------------------------------------- formats
def _fcfa(value):
    if value is None or value == '':
        return ''
    return f"{Decimal(str(value)):,.0f}".replace(',', ' ')


def _cell_text(value, money=False):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.strftime('%d/%m/%Y')
    if money or isinstance(value, Decimal):
        return _fcfa(value)
    return str(value)


def to_csv(report):
    """CSV pour Excel en français : séparateur « ; », UTF-8 avec BOM, montants sans séparateur de milliers"""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=';', quoting=csv.QUOTE_MINIMAL)
    writer.writerow(report.columns)

    def raw(value):
        if value is None:
            return ''
        if isinstance(value, (date, datetime)):
            return value.strftime('%d/%m/%Y')
        if isinstance(value, Decimal):
            return str(value.quantize(Decimal('1')) if value == value.to_integral() else value).replace('.', ',')
        return str(value)

    for row in report.rows:
        writer.writerow([raw(v) for v in row])
    if report.totals:
        writer.writerow([raw(v) for v in report.totals])
    return ('﻿' + buffer.getvalue()).encode('utf-8')


def to_excel(report, tontine_name):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = REPORT_SHORT_NAMES.get(report.kind, 'Rapport')[:30]
    ws.append([f"{tontine_name} - {report.title}"])
    ws['A1'].font = Font(bold=True, size=14, color='14532D')
    ws.append([f"Période {report.period}"])
    ws.append([])
    ws.append(report.columns)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='14532D')
        cell.alignment = Alignment(horizontal='center')
    for row in report.rows:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
    if report.totals:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in report.totals])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='FAF2D3')
    for idx in report.money_columns:
        for (cell,) in ws.iter_rows(min_row=header_row + 1, min_col=idx + 1, max_col=idx + 1):
            cell.number_format = '#,##0'
    for i, _ in enumerate(report.columns, 1):
        for (cell,) in ws.iter_rows(min_row=header_row + 1, min_col=i, max_col=i):
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = 'DD/MM/YYYY'
        ws.column_dimensions[get_column_letter(i)].width = 16 if i != 2 else 26
    if report.summary:
        ws.append([])
        ws.append(['Synthèse'])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for label, value in report.summary:
            ws.append([label.strip(), float(value)])
            ws.cell(ws.max_row, 2).number_format = '#,##0'
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(report, tontine_name, app_name, generated_by):
    """PDF A4 prêt à imprimer : en-tête avec logo, tableau, totaux, synthèse, pagination"""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    green, gold = colors.HexColor(GREEN), colors.HexColor(GOLD)
    wide = len(report.columns) > 6
    pagesize = landscape(A4) if wide else A4
    generated_at = datetime.now().strftime('%d/%m/%Y à %H:%M')

    class NumberedCanvas(rl_canvas.Canvas):
        """Ajoute « page X / Y » une fois le nombre total de pages connu"""
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved = []

        def showPage(self):
            self._saved.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._saved)
            for state in self._saved:
                self.__dict__.update(state)
                width, height = self._pagesize
                self.setFont('Helvetica', 8)
                self.setFillColor(colors.grey)
                self.drawString(15 * mm, 10 * mm, f"{app_name} · {tontine_name} · généré le {generated_at} par {generated_by}")
                self.drawRightString(width - 15 * mm, 10 * mm, f"Page {self._pageNumber} / {total}")
                super().showPage()
            super().save()

    def header(canv, doc):
        width, height = doc.pagesize
        top = height - 15 * mm
        # Logo : pièce verte cerclée d'or avec un G
        cx, cy, r = 15 * mm + 7 * mm, top - 7 * mm, 7 * mm
        canv.setFillColor(green)
        canv.circle(cx, cy, r, stroke=0, fill=1)
        canv.setStrokeColor(gold)
        canv.setLineWidth(1.2)
        canv.circle(cx, cy, r * 0.8, stroke=1, fill=0)
        canv.setFillColor(colors.white)
        canv.setFont('Helvetica-Bold', 13)
        canv.drawCentredString(cx, cy - 4.5, 'G')
        canv.setFillColor(green)
        canv.setFont('Helvetica-Bold', 15)
        canv.drawString(32 * mm, top - 5 * mm, tontine_name)
        canv.setFillColor(colors.black)
        canv.setFont('Helvetica', 10)
        canv.drawString(32 * mm, top - 10.5 * mm, f"{report.title} — période {report.period}")
        canv.setStrokeColor(gold)
        canv.setLineWidth(1.5)
        canv.line(15 * mm, top - 15 * mm, width - 15 * mm, top - 15 * mm)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=pagesize, leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=35 * mm, bottomMargin=18 * mm, title=f"{report.title} - {tontine_name}",
                            author=app_name)
    cell_style = ParagraphStyle('cell', fontName='Helvetica', fontSize=8, leading=10, alignment=TA_LEFT)
    story = []

    data = [report.columns]
    for row in report.rows:
        data.append([_cell_text(v, i in report.money_columns) if i in report.money_columns or not isinstance(v, str)
                     else Paragraph(v.replace('&', '&amp;').replace('<', '&lt;'), cell_style)
                     for i, v in enumerate(row)])
    if report.totals:
        data.append([_cell_text(v, i in report.money_columns) for i, v in enumerate(report.totals)])
    if not report.rows:
        data.append(['Aucune donnée sur cette période'] + [''] * (len(report.columns) - 1))

    available = doc.width
    weights = [1.0] * len(report.columns)
    weights[1] = 1.8  # colonne « Membre »
    for i, col in enumerate(report.columns):
        if col in ('Description', 'Opération', 'Rubrique', 'Motif'):
            weights[i] = 1.8
    col_widths = [available * w / sum(weights) for w in weights]

    table = LongTable(data, colWidths=col_widths, repeatRows=1)
    style = [
        ('FONT', (0, 0), (-1, 0), 'Helvetica-Bold', 8),
        ('FONT', (0, 1), (-1, -1), 'Helvetica', 8),
        ('BACKGROUND', (0, 0), (-1, 0), green),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#d9d6c8')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]
    for i in report.money_columns:
        style.append(('ALIGN', (i, 1), (i, -1), 'RIGHT'))
    for r in range(1, len(data) - (1 if report.totals else 0)):
        if r % 2 == 0:
            style.append(('BACKGROUND', (0, r), (-1, r), colors.HexColor(GREEN_LIGHT)))
    if report.totals and report.rows:
        style += [('BACKGROUND', (0, -1), (-1, -1), colors.HexColor(GOLD_LIGHT)),
                  ('FONT', (0, -1), (-1, -1), 'Helvetica-Bold', 8),
                  ('LINEABOVE', (0, -1), (-1, -1), 1, gold)]
    if not report.rows:
        style.append(('SPAN', (0, 1), (-1, 1)))
    table.setStyle(TableStyle(style))
    story.append(table)

    if report.summary:
        story.append(Spacer(1, 8 * mm))
        summary_data = [['Synthèse', 'Montant (FCFA)']] + [[label.strip(), _fcfa(value)] for label, value in report.summary]
        summary = Table(summary_data, colWidths=[70 * mm, 40 * mm], hAlign='LEFT')
        summary.setStyle(TableStyle([
            ('FONT', (0, 0), (-1, 0), 'Helvetica-Bold', 9),
            ('FONT', (0, 1), (-1, -1), 'Helvetica', 9),
            ('BACKGROUND', (0, 0), (-1, 0), green),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
            ('LINEBELOW', (0, 0), (-1, -1), 0.25, colors.HexColor('#d9d6c8')),
        ]))
        story.append(summary)

    doc.build(story, onFirstPage=header, onLaterPages=header, canvasmaker=NumberedCanvas)
    return buffer.getvalue()
